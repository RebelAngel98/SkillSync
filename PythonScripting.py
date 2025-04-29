import pandas as pd
import openai
import os
from sklearn.metrics.pairwise import cosine_similarity
from openpyxl import load_workbook


# openai.verify_ssl_certs = False  # or use a cert file path if available

# === CONFIGURATION ===
EXCEL_FILE = "Lightcast Skills.xlsx"
SOURCE_SHEET = "Lightcast_Skills_v9.29"
OUTPUT_SHEET = "Skill_Job_Matches"
API_KEY = "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"  # Or load from os.environ


# Example job titles list (in a real scenario, use a CSV or API)
job_titles = [
    "Software Engineer",
    "Full Stack Developer",
    "Frontend Developer",
    "Backend Developer",
    "DevOps Engineer",
    "Data Scientist",
    "Machine Learning Engineer",
    "AI Engineer",
    "Cybersecurity Analyst",
    "Network Engineer",
    "Database Administrator",
    "Cloud Engineer",
    "System Administrator",
    "IT Support Specialist",
    "Web Developer",
    "UI/UX Designer",
    "Product Manager (Tech)",
    "Business Intelligence Analyst",
    "IT Manager",
    "Software Architect",
    "Technical Lead",
    "Solutions Architect",
    "IT Consultant",
    "Security Engineer",
    "Automation Engineer",
    "Cloud Architect",
    "Mobile App Developer",
    "Game Developer",
    "Quality Assurance (QA) Engineer",
    "Systems Analyst",
    "IT Director",
    "Cloud Solutions Engineer",
    "Enterprise Architect",
    "SAP Consultant",
    "Blockchain Developer",
    "Data Engineer",
    "IT Specialist",
    "Application Support Analyst",
    "Technology Analyst",
    "Web Administrator",
    "Site Reliability Engineer",
    "Information Systems Manager",
    "Technical Writer",
    "Business Systems Analyst",
    "IT Operations Manager",
    "Database Analyst",
    "Network Administrator",
    "Cloud Consultant",
    "Cybersecurity Consultant",
    "IT Security Specialist",
    "Blockchain Architect",
    "Data Privacy Officer",
    "Technology Consultant",
    "Virtualization Engineer",
    "Infrastructure Engineer",
    "Penetration Tester",
    "Information Technology Analyst"
]

# === SETUP OPENAI ===
openai.api_key = API_KEY

def get_embedding(text):
    response = openai.Embedding.create(
        input=text,
        model="text-embedding-3-small"
    )
    return response['data'][0]['embedding']

def main():
    # Load the Excel sheet
    df = df = pd.read_excel(r"Lightcast Skills.xlsx", sheet_name="Lightcast_Skills_v9.29", engine="openpyxl")
    skills = df['NAME'].dropna().unique()

    print("Generating embeddings for skills...")
    skill_embeddings = {skill: get_embedding(skill) for skill in skills}

    print("Generating embeddings for job titles...")
    job_embeddings = {title: get_embedding(title) for title in job_titles}

    # Match skills to job titles using cosine similarity
    matches = []
    for skill, s_emb in skill_embeddings.items():
        sims = []
        for title, j_emb in job_embeddings.items():
            score = cosine_similarity([s_emb], [j_emb])[0][0]
            sims.append((title, score))
        sims.sort(key=lambda x: x[1], reverse=True)
        top_titles = [f"{title} ({round(score, 2)})" for title, score in sims[:5]]
        matches.append({'skill': skill, 'Top Matching Job Titles': "; ".join(top_titles)})

    result_df = pd.DataFrame(matches)

    # Write to a new sheet in the same Excel file
    print("Writing results to Excel...")
    book = load_workbook(EXCEL_FILE)
    with pd.ExcelWriter(EXCEL_FILE, engine='openpyxl', mode='a') as writer:
        writer.book = book
        if OUTPUT_SHEET in writer.book.sheetnames:
            del writer.book[OUTPUT_SHEET]  # Remove if already exists
        result_df.to_excel(writer, sheet_name=OUTPUT_SHEET, index=False)

    print("Done! Check the new tab in your Excel file.")

if __name__ == "__main__":
    main()
